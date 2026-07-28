import serial
import serial.tools.list_ports
import csv
import openpyxl
from openpyxl.chart import LineChart, Reference
from datetime import datetime

# Auto detect Arduino port
ports = list(serial.tools.list_ports.comports())
arduino_port = None
for p in ports:
    if 'Arduino' in p.description or 'CH340' in p.description or 'USB Serial' in p.description:
        arduino_port = p.device
        break

if arduino_port is None:
    print("Arduino not found! Make sure it's plugged in.")
    input("Press Enter to exit...")
else:
    print(f"Found Arduino on {arduino_port}")
    ser = serial.Serial(arduino_port, 9600)
    filename = f"pressure_data_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    times = []
    s1_data = []
    s2_data = []
    s3_data = []

    print(f"Saving to {filename}... Press Ctrl+C to stop.")

    try:
        while True:
            line = ser.readline().decode('utf-8').strip()
            print(line)
            data = {}
            for part in line.split(','):
                if ':' in part:
                    key, value = part.split(':')
                    data[key.strip()] = value.strip()
            if 'Sensor1' in data and 'Sensor2' in data and 'Sensor3' in data:
                t = datetime.now().strftime('%H:%M:%S')
                times.append(t)
                s1_data.append(int(data['Sensor1']))
                s2_data.append(int(data['Sensor2']))
                s3_data.append(int(data['Sensor3']))
    except KeyboardInterrupt:
        print("\nStopped! Creating Excel file with chart...")

    # Create Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    # Write headers
    ws.append(['Time', 'Sensor 1', 'Sensor 2', 'Sensor 3'])

    # Write data
    for i in range(len(times)):
        ws.append([times[i], s1_data[i], s2_data[i], s3_data[i]])

    # Create line chart
    chart = LineChart()
    chart.title = "Pressure Sensor Readings"
    chart.style = 10
    chart.y_axis.title = "Pressure"
    chart.x_axis.title = "Time"

    data_ref = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=len(times)+1)
    chart.add_data(data_ref, titles_from_data=True)

    # Add chart to new sheet
    ws2 = wb.create_sheet("Chart")
    ws2.add_chart(chart, "A1")

    # Save file
    filepath = f"{__import__('os').path.expanduser('~')}\\Desktop\\{filename}"
    wb.save(filepath)
    print(f"Excel file saved to Desktop: {filename}")
    input("Press Enter to exit...")